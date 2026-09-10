#!/usr/bin/env python3
"""Render the seed corpus to files, with ground truth alongside.

Documents are laid out as PDFs, spreadsheets and a drawing, then the ones marked
with a scan profile are degraded into image-only PDFs. Because every file starts
from a known string, the ground truth written next to it is exact — which is
what makes OCR evaluation a measurement rather than an impression.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT / "apps" / "api" / "src"))
sys.path.insert(0, str(Path(__file__).parent))

from content import SEED_DOCUMENTS, SeedDocument  # noqa: E402
from degrade import PROFILES, pdf_to_scanned_pdf  # noqa: E402

PAGE_WIDTH, PAGE_HEIGHT = 595, 842  # A4 in points


def render_pdf(document: SeedDocument, target: Path) -> None:
    """Lay out a document as a native PDF."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    target.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body", parent=styles["BodyText"], fontSize=10, leading=14, alignment=TA_JUSTIFY
    )
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=13, spaceBefore=10, spaceAfter=6)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=11, spaceBefore=8, spaceAfter=4)
    title = ParagraphStyle("DocTitle", parent=styles["Title"], fontSize=15, spaceAfter=14)

    def header_footer(canvas, doc) -> None:  # type: ignore[no-untyped-def]
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        # A classification banner, as a real controlled document carries.
        canvas.drawString(20 * mm, A4[1] - 12 * mm, f"MRPL — {document.classification.upper()}")
        canvas.drawRightString(A4[0] - 20 * mm, A4[1] - 12 * mm, document.doc_id)
        canvas.drawCentredString(A4[0] / 2, 12 * mm, f"Page {doc.page}")
        canvas.restoreState()

    story: list = [Paragraph(document.title, title)]
    for section in document.sections:
        story.append(Paragraph(section.heading, h1 if section.level == 1 else h2))
        if section.body:
            story.append(Paragraph(section.body, body))
        if section.table:
            table = Table(section.table, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ]
                )
            )
            story.extend([Spacer(1, 4), table, Spacer(1, 6)])

    SimpleDocTemplate(
        str(target),
        pagesize=A4,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        title=document.title,
        author="MRPL",
    ).build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def render_tag_registry(target: Path) -> list[str]:
    """The equipment tag master, used to cross-check what a VLM reads."""
    from openpyxl import Workbook

    rows = [
        ("V-1201", "Pressure Vessel", "Crude Feed Surge Drum", "PID-CDU-001", "2029-03-06"),
        ("V-1202", "Pressure Vessel", "Reflux Accumulator", "PID-CDU-001", "2028-11-20"),
        ("P-101A", "Centrifugal Pump", "Crude Charge Pump A", "PID-CDU-002", "2029-01-14"),
        ("P-101B", "Centrifugal Pump", "Crude Charge Pump B", "PID-CDU-002", "2029-02-02"),
        ("E-2405", "Heat Exchanger", "Crude Preheat Exchanger", "PID-CDU-003", "2029-02-18"),
        ("HV-1205", "Control Valve", "V-1201 Feed Isolation", "PID-CDU-001", "2028-09-04"),
        ("XV-1207", "Shutdown Valve", "V-1201 Flare Isolation", "PID-CDU-001", "2028-09-04"),
        ("PSV-1207", "Relief Valve", "V-1201 Overpressure Protection", "PID-CDU-001", "2029-03-22"),
        ("PT-3402", "Pressure Transmitter", "V-1201 Vessel Pressure", "PID-CDU-001", "2029-01-30"),
        ("FI-1205", "Flow Indicator", "V-1201 Feed Flow", "PID-CDU-001", "2029-01-30"),
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tag Registry"
    sheet.append(["Tag", "Type", "Service", "P&ID Reference", "Last Inspection"])
    for row in rows:
        sheet.append(list(row))
    for column, width in zip("ABCDE", (12, 22, 34, 18, 16), strict=True):
        sheet.column_dimensions[column].width = width

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return [row[0] for row in rows]


def render_thickness_workbook(target: Path) -> None:
    """Thickness history as a spreadsheet, for the table-query path."""
    from content import THICKNESS_TABLE
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "V-1201 CML History"
    for row in THICKNESS_TABLE:
        sheet.append(row)

    readings = workbook.create_sheet("All Readings")
    readings.append(["Tag", "CML", "Year", "Thickness_mm"])
    for row in THICKNESS_TABLE[1:]:
        for year, value in zip(("2019", "2023", "2029"), row[2:5], strict=True):
            readings.append(["V-1201", row[0], year, float(value)])

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)


def render_pid(target: Path, tags: list[str]) -> dict[str, list[float]]:
    """Draw a simplified P&ID with tags at known positions.

    Known positions are the point: the vision model's output is checked against
    this, so tag extraction accuracy is measurable rather than eyeballed.
    """
    import cv2
    import numpy as np

    width, height = 2200, 1500
    canvas = np.full((height, width, 3), 255, np.uint8)
    black = (0, 0, 0)
    truth: dict[str, list[float]] = {}

    def label(text: str, x: int, y: int, scale: float = 0.7) -> None:
        cv2.putText(canvas, text, (x, y), cv2.FONT_HERSHEY_SIMPLEX, scale, black, 2, cv2.LINE_AA)
        (tw, th), _ = cv2.getTextSize(text, cv2.FONT_HERSHEY_SIMPLEX, scale, 2)
        truth[text] = [x / width, (y - th) / height, (x + tw) / width, y / height]

    cv2.rectangle(canvas, (40, 40), (width - 40, height - 40), black, 3)
    cv2.putText(
        canvas,
        "MRPL  CRUDE DISTILLATION UNIT   P&ID  PID-CDU-001",
        (70, 100),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.9,
        black,
        2,
        cv2.LINE_AA,
    )

    # Vessel
    cv2.rectangle(canvas, (900, 400), (1200, 900), black, 3)
    cv2.ellipse(canvas, (1050, 400), (150, 60), 0, 180, 360, black, 3)
    cv2.ellipse(canvas, (1050, 900), (150, 60), 0, 0, 180, black, 3)
    label("V-1201", 960, 660, 0.9)
    label("SURGE DRUM", 940, 700, 0.5)

    # Feed line and isolation valve
    cv2.line(canvas, (300, 650), (900, 650), black, 3)
    cv2.drawMarker(canvas, (600, 650), black, cv2.MARKER_TRIANGLE_UP, 40, 3)
    label("HV-1205", 540, 610)
    label("FI-1205", 380, 610)

    # Flare line and shutdown valve
    cv2.line(canvas, (1200, 500), (1850, 500), black, 3)
    cv2.drawMarker(canvas, (1500, 500), black, cv2.MARKER_TRIANGLE_UP, 40, 3)
    label("XV-1207", 1440, 460)
    label("TO FLARE", 1700, 460, 0.5)

    # Relief valve and instrument
    cv2.line(canvas, (1050, 340), (1050, 220), black, 3)
    cv2.circle(canvas, (1050, 190), 34, black, 3)
    label("PSV-1207", 1100, 200)
    cv2.circle(canvas, (820, 500), 34, black, 3)
    label("PT-3402", 700, 470)

    # Pumps
    for cx, tag in ((500, "P-101A"), (500 + 260, "P-101B")):
        cv2.circle(canvas, (cx, 1150), 55, black, 3)
        label(tag, cx - 55, 1250)
    cv2.line(canvas, (300, 1150), (445, 1150), black, 3)

    # Exchanger
    cv2.rectangle(canvas, (1450, 1050), (1800, 1250), black, 3)
    label("E-2405", 1560, 1160, 0.9)

    target.parent.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(target), canvas)
    return truth


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=REPO_ROOT / "data" / "seed" / "corpus")
    parser.add_argument("--truth", type=Path, default=REPO_ROOT / "data" / "seed" / "ground_truth")
    args = parser.parse_args()

    corpus, truth_dir = args.out, args.truth
    corpus.mkdir(parents=True, exist_ok=True)
    truth_dir.mkdir(parents=True, exist_ok=True)

    manifest: list[dict] = []
    page_truth: list[dict] = []

    for document in SEED_DOCUMENTS:
        native = corpus / document.doc_type / f"{document.doc_id}.pdf"
        render_pdf(document, native)

        final = native
        if document.scan_profile:
            scanned = corpus / document.doc_type / f"{document.doc_id}-scanned.pdf"
            pdf_to_scanned_pdf(native, scanned, PROFILES[document.scan_profile], seed=42)
            native.unlink()
            final = scanned

        manifest.append(
            {
                "doc_id": document.doc_id,
                "path": str(final.relative_to(corpus)),
                "title": document.title,
                "doc_type": document.doc_type,
                "classification": document.classification,
                "departments": document.departments,
                "tags": document.tags,
                "scan_profile": document.scan_profile or "native",
            }
        )
        page_truth.append({"doc_id": document.doc_id, "text": document.plain_text})
        print(f"  {document.doc_id:<18} {final.relative_to(corpus)}")

    tags = render_tag_registry(corpus / "tabular" / "TAG-REGISTRY.xlsx")
    render_thickness_workbook(corpus / "tabular" / "V-1201-THICKNESS.xlsx")
    manifest += [
        {
            "doc_id": "TAG-REGISTRY",
            "path": "tabular/TAG-REGISTRY.xlsx",
            "title": "MRPL Equipment Tag Registry",
            "doc_type": "tabular",
            "classification": "internal",
            "departments": ["inspection", "maintenance"],
            "tags": tags,
            "scan_profile": "native",
        },
        {
            "doc_id": "V-1201-THICKNESS",
            "path": "tabular/V-1201-THICKNESS.xlsx",
            "title": "V-1201 Thickness Measurement History",
            "doc_type": "tabular",
            "classification": "confidential",
            "departments": ["inspection"],
            "tags": ["V-1201"],
            "scan_profile": "native",
        },
    ]
    print("  TAG-REGISTRY       tabular/TAG-REGISTRY.xlsx")
    print("  V-1201-THICKNESS   tabular/V-1201-THICKNESS.xlsx")

    pid_truth = render_pid(corpus / "drawings" / "PID-CDU-001.png", tags)
    manifest.append(
        {
            "doc_id": "PID-CDU-001",
            "path": "drawings/PID-CDU-001.png",
            "title": "P&ID CDU-001 — V-1201 Surge Drum",
            "doc_type": "drawing",
            "classification": "confidential",
            "departments": ["inspection", "operations"],
            "tags": sorted(pid_truth),
            "scan_profile": "native",
        }
    )
    print("  PID-CDU-001        drawings/PID-CDU-001.png")

    (truth_dir / "doc_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    (truth_dir / "ocr_pages.jsonl").write_text(
        "\n".join(json.dumps(t) for t in page_truth), encoding="utf-8"
    )
    (truth_dir / "pid_tags.json").write_text(
        json.dumps({"drawing": "PID-CDU-001", "tags": pid_truth}, indent=2), encoding="utf-8"
    )
    print(f"\n{len(manifest)} documents, ground truth in {truth_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

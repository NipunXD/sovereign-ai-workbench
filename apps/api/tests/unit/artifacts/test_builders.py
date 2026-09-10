"""Artifact builders.

Round-trip tests: build the file, reopen it with the library a recipient would
use, and assert the structure survived. A builder that produces a corrupt file
still returns bytes and still passes a size check.
"""

from __future__ import annotations

import io

from workbench.artifacts.docx_builder import Block, DocxBuilder, DocxSpec
from workbench.artifacts.provenance import Provenance
from workbench.artifacts.xlsx_builder import Column, Sheet, XlsxBuilder, XlsxSpec


def test_word_report_round_trips() -> None:
    from docx import Document

    artifact = DocxBuilder().build(
        DocxSpec(
            title="V-1201 Assessment",
            classification="confidential",
            blocks=[
                Block(type="heading", text="1. Findings", level=1),
                Block(type="paragraph", text="CML-04 is the controlling location."),
                Block(type="bullets", items=["Rate 0.55 mm/yr", "Remaining life 2.18 yr"]),
                Block(type="table", rows=[["CML", "2029"], ["CML-04", "9.20"]]),
            ],
        ),
        Provenance(generated_by="senior"),
    )
    document = Document(io.BytesIO(artifact.data))
    text = "\n".join(p.text for p in document.paragraphs)

    assert "1. Findings" in text
    assert "CML-04 is the controlling location." in text
    assert "Rate 0.55 mm/yr" in text
    # The content table plus the provenance table.
    assert len(document.tables) == 2
    assert document.tables[0].cell(0, 0).text == "CML"


def test_the_classification_appears_in_the_page_header() -> None:
    """A controlled document that is marked only on its cover is one photocopied
    page away from being unmarked."""
    from docx import Document

    artifact = DocxBuilder().build(
        DocxSpec(title="Restricted Report", classification="restricted"),
        Provenance(),
    )
    document = Document(io.BytesIO(artifact.data))
    header = document.sections[0].header.paragraphs[0].text
    assert "RESTRICTED" in header


def test_spreadsheet_keeps_formulas_live() -> None:
    """An engineer will edit a reading and expect the result to update. A sheet
    of constants looks identical until someone does, and then it is silently
    wrong."""
    from openpyxl import load_workbook

    artifact = XlsxBuilder().build(
        XlsxSpec(
            title="Corrosion",
            sheets=[
                Sheet(
                    name="Data",
                    columns=[
                        Column(header="CML", kind="text"),
                        Column(header="2023", kind="number"),
                        Column(header="2029", kind="number"),
                        Column(header="Rate", kind="formula"),
                    ],
                    rows=[["CML-04", "12.50", "9.20", "=(B{row}-C{row})/6"]],
                )
            ],
        ),
        Provenance(),
    )
    sheet = load_workbook(io.BytesIO(artifact.data))["Data"]
    assert sheet["D2"].value == "=(B2-C2)/6"
    # Numbers must be stored as numbers, or they will not sort, chart, or feed
    # a formula.
    assert isinstance(sheet["B2"].value, float)
    assert sheet["B2"].value == 12.5


def test_every_workbook_carries_a_provenance_sheet() -> None:
    from openpyxl import load_workbook

    artifact = XlsxBuilder().build(
        XlsxSpec(title="Anything", sheets=[Sheet(name="Data")]), Provenance(generated_by="senior")
    )
    workbook = load_workbook(io.BytesIO(artifact.data))
    assert "Provenance" in workbook.sheetnames


def test_sheet_names_are_made_legal() -> None:
    """Excel rejects several characters and caps names at 31 characters; a
    model-supplied name will eventually hit both."""
    from openpyxl import load_workbook

    artifact = XlsxBuilder().build(
        XlsxSpec(
            title="x",
            sheets=[Sheet(name="Readings: V-1201/CML [bottom shell] survey 2029")],
        ),
        Provenance(),
    )
    name = load_workbook(io.BytesIO(artifact.data)).sheetnames[0]
    assert len(name) <= 31
    assert not set(name) & set("[]:*?/\\")


def test_identical_inputs_produce_an_identical_digest() -> None:
    """Content addressing depends on this: regenerating the same report must
    not accumulate copies."""
    spec = DocxSpec(title="Stable", blocks=[Block(type="paragraph", text="same")])
    first = DocxBuilder().build(spec, Provenance(run_id="r1"))
    second = DocxBuilder().build(spec, Provenance(run_id="r1"))
    assert first.sha256 == second.sha256


def test_flat_tool_input_maps_onto_the_builder_spec() -> None:
    """The tool's schema is flattened for the model; the mapping must preserve
    meaning."""
    from workbench.tools.artifacts import XlsxInput, _to_xlsx_spec

    spec = _to_xlsx_spec(
        XlsxInput(
            title="Readings",
            headers=["CML", "2023", "2029"],
            rows=[["CML-04", "12.50", "9.20"]],
            numeric_headers=["2023", "2029"],
            formulas={"Rate": "=(B{row}-C{row})/6"},
            chart_of=["2023", "2029"],
        )
    )
    sheet = spec.sheets[0]
    assert [c.header for c in sheet.columns] == ["CML", "2023", "2029", "Rate"]
    assert [c.kind for c in sheet.columns] == ["text", "number", "number", "formula"]
    # The formula is appended to every row.
    assert sheet.rows[0][-1] == "=(B{row}-C{row})/6"
    assert sheet.chart is not None and sheet.chart.value_columns == [2, 3]

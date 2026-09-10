"""PowerPoint deck generation.

Round-trip tests, in the same spirit as the Word and Excel ones: build the
deck, reopen it with the library a recipient's tooling would use, and assert
the structure survived.
"""

from __future__ import annotations

import io
import struct
import zlib

import pytest

from workbench.artifacts.pptx_builder import (
    MAX_BULLETS_PER_SLIDE,
    PptxBuilder,
    PptxSpec,
    Slide,
)
from workbench.artifacts.provenance import Provenance
from workbench.rag.citations import Citation


@pytest.fixture
def png_bytes() -> bytes:
    """A real 1x1 PNG, built rather than pasted as a base64 blob.

    python-pptx parses the header to size the picture, so a placeholder that is
    not actually a PNG fails in a way unrelated to what the test is checking.
    """

    def chunk(tag: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload))
            + tag
            + payload
            + struct.pack(">I", zlib.crc32(tag + payload) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixel = zlib.compress(b"\x00\xff\xff\xff")
    return (
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", pixel) + chunk(b"IEND", b"")
    )


def _open(data: bytes):
    from pptx import Presentation

    return Presentation(io.BytesIO(data))


def _text(slide) -> str:
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def test_deck_round_trips() -> None:
    artifact = PptxBuilder().build(
        PptxSpec(
            title="V-1201 Turnaround Briefing",
            subtitle="March 2029",
            classification="confidential",
            slides=[
                Slide(type="section", title="Condition Summary"),
                Slide(
                    type="bullets",
                    title="Findings",
                    bullets=["CML-04 is controlling", "Rate 0.55 mm/yr"],
                    notes="INSP-2029-V1201 p.2",
                ),
            ],
        ),
        Provenance(generated_by="senior"),
    )
    assert artifact.filename == "V-1201-Turnaround-Briefing.pptx"
    assert artifact.mime.endswith("presentationml.presentation")

    deck = _open(artifact.data)
    body = "\n".join(_text(slide) for slide in deck.slides)
    assert "V-1201 Turnaround Briefing" in body
    assert "Condition Summary" in body
    assert "CML-04 is controlling" in body
    assert "INSP-2029-V1201 p.2" in body


def test_every_slide_carries_the_classification() -> None:
    """A deck is the format most likely to be broken up and re-pasted, so the
    marking cannot live only on the cover."""
    artifact = PptxBuilder().build(
        PptxSpec(
            title="Restricted Briefing",
            classification="restricted",
            slides=[
                Slide(type="bullets", title="One", bullets=["a"]),
                Slide(type="table", title="Two", rows=[["h"], ["v"]]),
                Slide(type="quote", title="Three"),
            ],
        ),
        Provenance(),
    )
    deck = _open(artifact.data)
    assert len(deck.slides) >= 6  # cover + 3 + provenance + sources
    for slide in deck.slides:
        assert "MRPL — RESTRICTED" in _text(slide)


def test_long_bullet_lists_continue_rather_than_truncate() -> None:
    count = MAX_BULLETS_PER_SLIDE * 2 + 1
    artifact = PptxBuilder().build(
        PptxSpec(
            title="Deck",
            slides=[
                Slide(
                    type="bullets",
                    title="Findings",
                    bullets=[f"point {i}" for i in range(count)],
                )
            ],
        ),
        Provenance(),
    )
    deck = _open(artifact.data)
    body = "\n".join(_text(slide) for slide in deck.slides)
    # Nothing is dropped, and the overflow is marked as a continuation.
    for i in range(count):
        assert f"point {i}" in body
    assert "Findings (cont.)" in body


def test_an_oversized_table_says_that_it_was_cut() -> None:
    """Silently showing the first eleven of forty rows is worse than a small
    table, because nobody in the room knows the rest exist."""
    rows = [["CML", "2029"]] + [[f"CML-{i:02d}", "9.20"] for i in range(1, 40)]
    artifact = PptxBuilder().build(
        PptxSpec(title="Deck", slides=[Slide(type="table", title="Data", rows=rows)]),
        Provenance(),
    )
    deck = _open(artifact.data)
    body = "\n".join(_text(slide) for slide in deck.slides)
    assert "further rows omitted" in body


def test_a_missing_image_is_stated_on_the_slide() -> None:
    """A chart that vanished between the sandbox and the deck must not leave a
    blank frame the audience discovers first."""
    artifact = PptxBuilder(images={}).build(
        PptxSpec(
            title="Deck",
            slides=[Slide(type="image", title="Trend", image_ref="trend.png")],
        ),
        Provenance(),
    )
    body = "\n".join(_text(slide) for slide in _open(artifact.data).slides)
    assert "trend.png" in body
    assert "not produced" in body


def test_an_embedded_image_becomes_a_picture(png_bytes: bytes) -> None:
    artifact = PptxBuilder(images={"trend.png": png_bytes}).build(
        PptxSpec(
            title="Deck",
            slides=[Slide(type="image", title="Trend", image_ref="trend.png")],
        ),
        Provenance(),
    )
    deck = _open(artifact.data)
    pictures = [
        shape
        for slide in deck.slides
        for shape in slide.shapes
        if shape.shape_type is not None and "PICTURE" in str(shape.shape_type)
    ]
    assert pictures


def test_provenance_travels_inside_the_deck() -> None:
    """The deck is forwarded on its own, so the record has to be in the file."""
    provenance = Provenance(
        run_id="run_1",
        generated_by="a.kumar",
        models={"reasoning": "qwen/qwen3-8b"},
        citations=[
            Citation(
                n=1,
                chunk_id="c1",
                doc_id="INSP-2029-V1201",
                doc_title="Inspection Report V-1201",
                doc_type="inspection",
                page_no=2,
                confidence=0.61,
            )
        ],
    )
    artifact = PptxBuilder().build(PptxSpec(title="Deck"), provenance)
    body = "\n".join(_text(slide) for slide in _open(artifact.data).slides)

    assert "a.kumar" in body
    assert "run_1" in body
    assert "Inspection Report V-1201" in body
    # A figure read from a bad photocopy is flagged, not quietly cited.
    assert "61%" in body or "confidence" in body


def test_an_ungrounded_deck_says_so() -> None:
    artifact = PptxBuilder().build(PptxSpec(title="Deck"), Provenance())
    body = "\n".join(_text(slide) for slide in _open(artifact.data).slides)
    assert "No documents were cited" in body

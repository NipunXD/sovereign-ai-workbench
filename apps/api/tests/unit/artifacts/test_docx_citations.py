"""Inline citations in a generated Word document.

The body used to carry no markers at all — sources were listed on the
provenance page, and a reader at "9.20 mm" had no way to know which of the
eight listed documents that figure came from. The chat answer already cites
by the number of the source shown to the model; the document's arguments are
generated against the same numbered list, so the same [n] resolves the same
way.
"""

from __future__ import annotations

import io

from workbench.artifacts.docx_builder import Block, DocxBuilder, DocxSpec
from workbench.artifacts.provenance import Provenance
from workbench.core.citation import Citation


def _sources() -> list[Citation]:
    return [
        Citation(
            n=1,
            chunk_id="c1",
            doc_id="d1",
            doc_title="Inspection Report V-1201 — March 2029",
            page_no=2,
            section_path=["4. Findings"],
            snippet="CML-04 on the bottom shell measured 9.20 mm",
            confidence=0.61,
        ),
        Citation(
            n=2,
            chunk_id="c2",
            doc_id="d2",
            doc_title="SOP-4412 Vessel Shutdown and Depressurisation Procedure",
            page_no=1,
            section_path=["7. Retirement Criteria"],
            snippet="The minimum required shell thickness for V-1201 is 8.0 mm",
        ),
        Citation(
            n=3,
            chunk_id="c3",
            doc_id="d3",
            doc_title="Maintenance Activity Log — Q1 2029",
            page_no=1,
        ),
    ]


def _build(blocks: list[Block]) -> tuple[bytes, dict]:
    artifact = DocxBuilder().build(
        DocxSpec(title="Survey", blocks=blocks), Provenance(citations=_sources())
    )
    return artifact.data, artifact.meta


def _document(data: bytes):
    from docx import Document

    return Document(io.BytesIO(data))


def test_a_marker_becomes_a_superscript_reference() -> None:
    data, meta = _build(
        [
            Block(
                type="paragraph",
                text="CML-04 measured 9.20 mm [1] against a minimum of 8.0 mm [2].",
            )
        ]
    )
    document = _document(data)

    paragraph = next(p for p in document.paragraphs if "CML-04" in p.text)
    assert paragraph.text == "CML-04 measured 9.20 mm [1] against a minimum of 8.0 mm [2]."
    superscripts = [r.text for r in paragraph.runs if r.font.superscript]
    assert superscripts == ["[1]", "[2]"]
    assert meta["cited_sources"] == [1, 2]


def test_markers_survive_inside_table_cells_and_bullets() -> None:
    data, meta = _build(
        [
            Block(type="table", rows=[["CML", "2029 (mm)"], ["CML-04", "9.20 [1]"]]),
            Block(type="bullets", items=["Retire at 8.0 mm [2]"]),
        ]
    )
    document = _document(data)

    cell = document.tables[0].cell(1, 1)
    assert cell.text == "9.20 [1]"
    assert any(r.font.superscript for p in cell.paragraphs for r in p.runs)
    bullet = next(p for p in document.paragraphs if "Retire" in p.text)
    assert [r.text for r in bullet.runs if r.font.superscript] == ["[2]"]
    assert meta["cited_sources"] == [1, 2]


def test_a_marker_with_no_source_behind_it_is_removed_and_reported() -> None:
    """A reference to nothing is worse than no reference — same rule as the chat."""
    data, meta = _build([Block(type="paragraph", text="The rate is 0.55 mm/yr [7].")])
    document = _document(data)

    paragraph = next(p for p in document.paragraphs if "0.55" in p.text)
    assert paragraph.text == "The rate is 0.55 mm/yr."
    assert meta["unresolved_markers"] == ["[7]"]
    text = "\n".join(p.text for p in document.paragraphs)
    assert "referring to no source ([7]) were removed" in text


def test_doubled_brackets_are_normalised() -> None:
    data, _ = _build([Block(type="paragraph", text="Highest loss at CML-04 [[1]].")])
    paragraph = next(p for p in _document(data).paragraphs if "Highest" in p.text)
    assert paragraph.text == "Highest loss at CML-04 [1]."


def test_the_references_list_is_numbered_to_match_the_markers() -> None:
    """The reader at "[2]" must find "2." and nothing else answering to it."""
    data, _ = _build([Block(type="paragraph", text="Minimum thickness is 8.0 mm [2].")])
    document = _document(data)
    text = "\n".join(p.text for p in document.paragraphs)

    sources = text[text.index("Sources") :]
    assert "1.  Inspection Report V-1201 — March 2029 — page 2 — 4. Findings" in sources
    assert "2.  SOP-4412 Vessel Shutdown and Depressurisation Procedure — page 1" in sources
    assert "3.  Maintenance Activity Log — Q1 2029 — page 1" in sources
    # Order is the source order the model was shown, never re-sorted.
    assert (
        sources.index("1.  Inspection")
        < sources.index("2.  SOP-4412")
        < sources.index("3.  Maintenance")
    )


def test_sources_read_but_not_cited_are_marked_as_such() -> None:
    data, _ = _build([Block(type="paragraph", text="Minimum thickness is 8.0 mm [2].")])
    text = "\n".join(p.text for p in _document(data).paragraphs)

    for line in text.splitlines():
        if line.startswith("1.  Inspection") or line.startswith("3.  Maintenance"):
            assert "retrieved, not cited" in line
        if line.startswith("2.  SOP-4412"):
            assert "retrieved, not cited" not in line


def test_a_scanned_source_keeps_its_confidence_warning() -> None:
    data, _ = _build([Block(type="paragraph", text="9.20 mm [1].")])
    text = "\n".join(p.text for p in _document(data).paragraphs)
    line = next(row for row in text.splitlines() if row.startswith("1.  Inspection"))
    assert "61% confidence" in line
    assert "verify against the original" in line


def test_text_without_markers_is_unchanged() -> None:
    data, meta = _build(
        [Block(type="paragraph", text="Readings are in [mm] and [12.50, 9.20] were recorded.")]
    )
    paragraph = next(p for p in _document(data).paragraphs if "Readings" in p.text)
    # Neither a unit in brackets nor a bracketed list is a source marker.
    assert paragraph.text == "Readings are in [mm] and [12.50, 9.20] were recorded."
    assert meta["cited_sources"] == []
    assert meta["unresolved_markers"] == []

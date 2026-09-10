"""Word, PowerPoint and spreadsheet extraction.

Office formats carry their own structure — heading styles, slide layouts, sheet
and cell types — so this reads the object model rather than rendering to text
and guessing. Headings come from the style name, not from font size heuristics.

Spreadsheets are handled differently from prose on purpose. A maintenance log
with four thousand rows is useless as text in a context window; it is registered
as a queryable dataset and only a representative sample is indexed for search,
so the agent can *find* the sheet and then compute over it.
"""

from __future__ import annotations

import time
from pathlib import Path
from typing import Any

from workbench.core.logging import get_logger
from workbench.ingest.ir import (
    BBox,
    Block,
    BlockSource,
    BlockType,
    DatasetRef,
    DocumentIR,
    ExtractionReport,
    Page,
)

log = get_logger(__name__)

#: Rows serialised into the searchable text sample for a sheet. Enough to make
#: the sheet findable by its contents; the full data lives in the dataset.
SAMPLE_ROWS = 50


class DocxExtractor:
    name = "docx"

    def extract(
        self, path: Path, doc_id: str, *, filename: str | None = None, **_: Any
    ) -> DocumentIR:
        from docx import Document as DocxDocument

        started = time.perf_counter()
        document = DocxDocument(str(path))
        blocks: list[Block] = []
        order = 0

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            block_type = BlockType.PARAGRAPH
            attrs: dict[str, Any] = {}
            if style.startswith("heading"):
                block_type = BlockType.HEADING
                # "Heading 2" -> level 2. Style names are the document's own
                # structure, which beats inferring hierarchy from font size.
                attrs["level"] = int(style.replace("heading", "").strip() or 1)
            elif style.startswith("title"):
                block_type, attrs["level"] = BlockType.HEADING, 1
            elif style.startswith(("list", "bullet")):
                block_type = BlockType.LIST

            blocks.append(
                Block(
                    block_id=f"p{order}",
                    page_no=1,
                    type=block_type,
                    text=text,
                    source=BlockSource.OFFICE,
                    order=order,
                    attrs=attrs,
                )
            )
            order += 1

        for index, table in enumerate(document.tables):
            rows = [
                " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
                for row in table.rows
            ]
            rows = [r for r in rows if r.strip(" |")]
            if not rows:
                continue
            blocks.append(
                Block(
                    block_id=f"t{index}",
                    page_no=1,
                    type=BlockType.TABLE,
                    text="\n".join(rows),
                    source=BlockSource.OFFICE,
                    order=order,
                    attrs={"rows": len(rows), "columns": len(table.columns)},
                )
            )
            order += 1

        core = document.core_properties
        return DocumentIR(
            doc_id=doc_id,
            title=(core.title or Path(filename or path.name).stem).strip(),
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            # Word has no fixed pagination without rendering; the whole document
            # is one logical page and chunks carry section paths instead.
            page_count=1,
            pages=[Page(page_no=1, width=1240, height=1754, blocks=blocks)],
            metadata={"author": core.author, "created": str(core.created or "")},
            extraction_report=ExtractionReport(
                extractor=self.name,
                pages_processed=1,
                blocks_extracted=len(blocks),
                stage_timings_ms={"total": (time.perf_counter() - started) * 1000},
            ),
        )


class PptxExtractor:
    name = "pptx"

    def extract(
        self, path: Path, doc_id: str, *, filename: str | None = None, **_: Any
    ) -> DocumentIR:
        from pptx import Presentation
        from pptx.util import Emu

        started = time.perf_counter()
        presentation = Presentation(str(path))
        slide_width = Emu(presentation.slide_width or 9144000)
        slide_height = Emu(presentation.slide_height or 6858000)
        pages: list[Page] = []

        for index, slide in enumerate(presentation.slides, start=1):
            blocks: list[Block] = []
            order = 0
            title_shape = slide.shapes.title

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                is_title = title_shape is not None and shape == title_shape
                blocks.append(
                    Block(
                        block_id=f"s{index}b{order}",
                        page_no=index,
                        type=BlockType.HEADING if is_title else BlockType.PARAGRAPH,
                        text=text,
                        bbox=BBox.from_pixels(
                            (
                                float(shape.left or 0),
                                float(shape.top or 0),
                                float((shape.left or 0) + (shape.width or 0)),
                                float((shape.top or 0) + (shape.height or 0)),
                            ),
                            float(slide_width),
                            float(slide_height),
                        ),
                        source=BlockSource.OFFICE,
                        order=order,
                        attrs={"level": 1} if is_title else {},
                    )
                )
                order += 1

            if slide.has_notes_slide and (notes := slide.notes_slide.notes_text_frame.text.strip()):
                blocks.append(
                    Block(
                        block_id=f"s{index}notes",
                        page_no=index,
                        type=BlockType.CAPTION,
                        text=notes,
                        source=BlockSource.OFFICE,
                        order=order,
                        attrs={"speaker_notes": True},
                    )
                )

            pages.append(Page(page_no=index, width=1280, height=720, blocks=blocks))

        return DocumentIR(
            doc_id=doc_id,
            title=(presentation.core_properties.title or Path(filename or path.name).stem).strip(),
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            page_count=len(pages),
            pages=pages,
            extraction_report=ExtractionReport(
                extractor=self.name,
                pages_processed=len(pages),
                blocks_extracted=sum(len(p.blocks) for p in pages),
                stage_timings_ms={"total": (time.perf_counter() - started) * 1000},
            ),
        )


class SpreadsheetExtractor:
    """Excel and CSV.

    Produces both a searchable text sample and a registered dataset. The sample
    makes the sheet findable; the dataset is what the ``table.query`` tool runs
    aggregations against.
    """

    name = "spreadsheet"

    def __init__(self, dataset_dir: Path | None = None) -> None:
        self.dataset_dir = dataset_dir

    def extract(
        self, path: Path, doc_id: str, *, filename: str | None = None, **_: Any
    ) -> DocumentIR:
        import pandas as pd

        started = time.perf_counter()
        suffix = path.suffix.lower()

        if suffix in {".csv", ".tsv"}:
            separator = "\t" if suffix == ".tsv" else ","
            frames = {"data": pd.read_csv(path, sep=separator, dtype=str, keep_default_na=False)}
            mime = "text/csv"
        else:
            frames = pd.read_excel(path, sheet_name=None, dtype=str, na_filter=False)
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        pages: list[Page] = []
        datasets: list[DatasetRef] = []

        for index, (sheet_name, frame) in enumerate(frames.items(), start=1):
            frame = frame.dropna(how="all").dropna(axis=1, how="all")
            if frame.empty:
                continue

            columns = [str(c) for c in frame.columns]
            header = " | ".join(columns)
            sample = frame.head(SAMPLE_ROWS)
            rows = [
                " | ".join(str(value) for value in record)
                for record in sample.itertuples(index=False, name=None)
            ]

            blocks = [
                Block(
                    block_id=f"sheet{index}h",
                    page_no=index,
                    type=BlockType.HEADING,
                    text=f"Sheet: {sheet_name}",
                    source=BlockSource.OFFICE,
                    order=0,
                    attrs={"level": 1},
                ),
                Block(
                    block_id=f"sheet{index}t",
                    page_no=index,
                    type=BlockType.TABLE,
                    text="\n".join([header, *rows]),
                    source=BlockSource.OFFICE,
                    order=1,
                    attrs={
                        "sheet": sheet_name,
                        "total_rows": len(frame),
                        "sampled_rows": len(rows),
                        "columns": columns,
                    },
                ),
            ]
            if len(frame) > SAMPLE_ROWS:
                # Say so explicitly. A model shown 50 of 4,000 rows will
                # otherwise answer "the maximum is X" about the sample.
                blocks.append(
                    Block(
                        block_id=f"sheet{index}n",
                        page_no=index,
                        type=BlockType.CAPTION,
                        text=(
                            f"This is a sample of {len(rows)} rows from {len(frame)} total. "
                            f"Query the dataset for complete results rather than reading "
                            f"figures from this sample."
                        ),
                        source=BlockSource.OFFICE,
                        order=2,
                    )
                )

            pages.append(Page(page_no=index, width=1240, height=1754, blocks=blocks))

            dataset_path = ""
            if self.dataset_dir is not None:
                self.dataset_dir.mkdir(parents=True, exist_ok=True)
                dataset_path = str(self.dataset_dir / f"{doc_id}_{index}.parquet")
                try:
                    # Written through DuckDB rather than pandas.to_parquet:
                    # DuckDB has a parquet writer built in, whereas pandas needs
                    # pyarrow, which is a 100 MB dependency that would otherwise
                    # be carried into the air-gap bundle for this one call.
                    import duckdb

                    connection = duckdb.connect()
                    connection.register("sheet_data", frame)
                    connection.execute("COPY sheet_data TO ? (FORMAT PARQUET)", [dataset_path])
                    connection.close()
                except Exception as exc:
                    log.warning("dataset_write_failed", sheet=sheet_name, error=str(exc))
                    dataset_path = ""

            datasets.append(
                DatasetRef(
                    dataset_id=f"{doc_id}_ds{index}",
                    sheet_name=str(sheet_name),
                    row_count=len(frame),
                    columns=[{"name": c, "dtype": "string"} for c in columns],
                    path=dataset_path,
                )
            )

        return DocumentIR(
            doc_id=doc_id,
            title=Path(filename or path.name).stem,
            mime=mime,
            page_count=len(pages),
            pages=pages,
            datasets=datasets,
            extraction_report=ExtractionReport(
                extractor=self.name,
                pages_processed=len(pages),
                blocks_extracted=sum(len(p.blocks) for p in pages),
                stage_timings_ms={"total": (time.perf_counter() - started) * 1000},
            ),
        )


class TextExtractor:
    """Plain text, markdown and email."""

    name = "text"

    def extract(
        self, path: Path, doc_id: str, *, filename: str | None = None, **_: Any
    ) -> DocumentIR:
        import re

        started = time.perf_counter()
        raw = path.read_text(encoding="utf-8", errors="replace")
        mime = "text/plain"
        metadata: dict[str, Any] = {}

        if path.suffix.lower() == ".eml":
            raw, metadata = self._parse_email(path)
            mime = "message/rfc822"

        blocks: list[Block] = []
        order = 0
        for paragraph in re.split(r"\n\s*\n", raw):
            text = paragraph.strip()
            if not text:
                continue
            # Markdown headings are explicit structure; use them.
            heading = re.match(r"^(#{1,6})\s+(.*)$", text)
            if heading:
                blocks.append(
                    Block(
                        block_id=f"b{order}",
                        page_no=1,
                        type=BlockType.HEADING,
                        text=heading.group(2).strip(),
                        source=BlockSource.NATIVE,
                        order=order,
                        attrs={"level": len(heading.group(1))},
                    )
                )
            else:
                blocks.append(
                    Block(
                        block_id=f"b{order}",
                        page_no=1,
                        type=BlockType.PARAGRAPH,
                        text=text,
                        source=BlockSource.NATIVE,
                        order=order,
                    )
                )
            order += 1

        return DocumentIR(
            doc_id=doc_id,
            title=str(metadata.get("subject") or Path(filename or path.name).stem),
            mime=mime,
            page_count=1,
            pages=[Page(page_no=1, width=1240, height=1754, blocks=blocks)],
            metadata=metadata,
            extraction_report=ExtractionReport(
                extractor=self.name,
                pages_processed=1,
                blocks_extracted=len(blocks),
                stage_timings_ms={"total": (time.perf_counter() - started) * 1000},
            ),
        )

    @staticmethod
    def _parse_email(path: Path) -> tuple[str, dict[str, Any]]:
        """Flatten an email to its text body plus the headers worth keeping."""
        import email
        from email import policy

        message = email.message_from_bytes(path.read_bytes(), policy=policy.default)
        metadata = {
            "subject": str(message.get("subject", "")),
            "from": str(message.get("from", "")),
            "to": str(message.get("to", "")),
            "date": str(message.get("date", "")),
        }
        body = message.get_body(preferencelist=("plain", "html"))
        text = body.get_content() if body else ""
        header_lines = "\n".join(f"{k.title()}: {v}" for k, v in metadata.items() if v)
        return f"{header_lines}\n\n{text}", metadata

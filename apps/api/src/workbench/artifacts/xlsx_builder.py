"""Spreadsheet generation.

The important decision here is that calculated cells carry *formulas*, not
precomputed numbers. An engineer who opens a corrosion workbook will change a
thickness reading and expect the remaining life to update. A sheet full of
constants looks identical until someone edits it, at which point it quietly
becomes wrong — which is worse than obviously wrong.

The provenance sheet is a real sheet rather than a comment, so it survives being
emailed onward and cannot be missed by someone who opens the file at the tab
they care about.
"""

from __future__ import annotations

import io
from typing import Any, Literal

from pydantic import BaseModel, Field

from workbench.artifacts.base import ArtifactBytes, ArtifactKind
from workbench.artifacts.provenance import AI_NOTICE, Provenance
from workbench.core.hashing import digest_bytes


class Column(BaseModel):
    header: str
    #: "text" | "number" | "formula". A formula cell's value is an Excel
    #: expression with {row} substituted for the 1-based sheet row.
    kind: Literal["text", "number", "formula"] = "text"
    number_format: str = ""
    width: int = 0


class ChartSpec(BaseModel):
    type: Literal["line", "bar"] = "line"
    title: str = ""
    #: 1-based column indices within the sheet.
    categories_column: int = 1
    value_columns: list[int] = Field(default_factory=list)
    y_axis_title: str = ""


class Sheet(BaseModel):
    name: str
    columns: list[Column] = Field(default_factory=list)
    rows: list[list[Any]] = Field(default_factory=list)
    chart: ChartSpec | None = None
    freeze_header: bool = True
    notes: list[str] = Field(default_factory=list)


class XlsxSpec(BaseModel):
    title: str
    classification: str = "internal"
    sheets: list[Sheet] = Field(default_factory=list)


class XlsxBuilder:
    kind = ArtifactKind.XLSX

    def build(self, spec: XlsxSpec, provenance: Provenance) -> ArtifactBytes:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        workbook.remove(workbook.active)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="44546A")

        for sheet_spec in spec.sheets:
            # Excel rejects several characters in sheet names and caps them at
            # 31 chars; a model-supplied name will eventually hit both.
            sheet = workbook.create_sheet(_safe_sheet_name(sheet_spec.name))

            for index, column in enumerate(sheet_spec.columns, start=1):
                cell = sheet.cell(row=1, column=index, value=column.header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                letter = get_column_letter(index)
                sheet.column_dimensions[letter].width = column.width or max(
                    12, min(40, len(column.header) + 4)
                )

            for row_offset, row in enumerate(sheet_spec.rows):
                excel_row = row_offset + 2
                for column_offset, column in enumerate(sheet_spec.columns):
                    raw = row[column_offset] if column_offset < len(row) else None
                    cell = sheet.cell(row=excel_row, column=column_offset + 1)

                    if column.kind == "formula" and isinstance(raw, str) and raw:
                        # {row} lets a spec express "the value on this row"
                        # without knowing where the data starts.
                        cell.value = raw.replace("{row}", str(excel_row))
                    elif column.kind == "number":
                        cell.value = _as_number(raw)
                    else:
                        cell.value = raw

                    if column.number_format:
                        cell.number_format = column.number_format

            if sheet_spec.freeze_header and sheet_spec.rows:
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = (
                    f"A1:{get_column_letter(len(sheet_spec.columns))}{len(sheet_spec.rows) + 1}"
                )

            if sheet_spec.chart and sheet_spec.rows and sheet_spec.chart.value_columns:
                chart = LineChart() if sheet_spec.chart.type == "line" else BarChart()
                chart.title = sheet_spec.chart.title or sheet_spec.name
                chart.y_axis.title = sheet_spec.chart.y_axis_title
                last_row = len(sheet_spec.rows) + 1
                for column_index in sheet_spec.chart.value_columns:
                    chart.add_data(
                        Reference(sheet, min_col=column_index, min_row=1, max_row=last_row),
                        titles_from_data=True,
                    )
                chart.set_categories(
                    Reference(
                        sheet,
                        min_col=sheet_spec.chart.categories_column,
                        min_row=2,
                        max_row=last_row,
                    )
                )
                chart.height, chart.width = 8, 16
                sheet.add_chart(chart, f"{get_column_letter(len(sheet_spec.columns) + 2)}2")

            if sheet_spec.notes:
                start = len(sheet_spec.rows) + 3
                for offset, note in enumerate(sheet_spec.notes):
                    cell = sheet.cell(row=start + offset, column=1, value=note)
                    cell.font = Font(italic=True, size=9)

        self._provenance_sheet(workbook, spec, provenance)

        buffer = io.BytesIO()
        workbook.save(buffer)
        data = buffer.getvalue()
        return ArtifactBytes(
            data=data,
            kind=self.kind,
            filename=_filename(spec.title),
            sha256=digest_bytes(data),
            meta={"sheets": [s.name for s in spec.sheets], "classification": spec.classification},
        )

    @staticmethod
    def _provenance_sheet(workbook: Any, spec: XlsxSpec, provenance: Provenance) -> None:
        from openpyxl.styles import Alignment, Font

        sheet = workbook.create_sheet("Provenance")
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 96

        sheet["A1"] = spec.title
        sheet["A1"].font = Font(bold=True, size=13)
        sheet["A2"] = f"Classification: {spec.classification.upper()}"
        sheet["A2"].font = Font(bold=True)

        row = 4
        for label, value in provenance.lines():
            sheet.cell(row=row, column=1, value=label).font = Font(bold=True)
            sheet.cell(row=row, column=2, value=value)
            row += 1

        row += 1
        sheet.cell(row=row, column=1, value="Sources").font = Font(bold=True, size=11)
        row += 1
        source_lines = provenance.source_lines()
        for line in source_lines or ["No documents were cited."]:
            sheet.cell(row=row, column=2, value=line)
            row += 1

        row += 1
        notice = sheet.cell(row=row, column=2, value=AI_NOTICE)
        notice.font = Font(italic=True, size=9)
        notice.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 46


def _as_number(value: Any) -> Any:
    """Coerce to a real number so Excel treats it as one.

    A numeric reading stored as text sorts wrongly, will not chart, and silently
    breaks any formula referring to it.
    """
    if value is None or isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return value


def _safe_sheet_name(name: str) -> str:
    import re

    cleaned = re.sub(r"[\[\]:*?/\\]", "-", name).strip() or "Sheet"
    return cleaned[:31]


def _filename(title: str) -> str:
    import re

    slug = re.sub(r"[^A-Za-z0-9]+", "-", title).strip("-")[:60] or "workbook"
    return f"{slug}.xlsx"
